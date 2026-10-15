import io

import mammoth
from docx import Document
from docx.shared import Pt

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def text_to_docx_bytes(text: str, title: str = "Job Description") -> bytes:
    """Build a simple .docx from plain text and return it as bytes.

    Uses a plain bold title (not Word's styled "Heading 1") so the document
    looks clean and neutral. Each input line becomes its own paragraph.
    """
    document = Document()

    title_paragraph = document.add_paragraph()
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph("")  # spacer line under the title

    for line in (text or "").splitlines():
        document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_bytes_to_html(data: bytes) -> str:
    """Convert .docx bytes to HTML for in-browser preview via mammoth."""
    result = mammoth.convert_to_html(io.BytesIO(data))
    return result.value  # ignores result.messages (style warnings)


def docx_bytes_to_text(data: bytes) -> str:
    """Extract plain text from .docx bytes (for feeding to the model)."""
    result = mammoth.extract_raw_text(io.BytesIO(data))
    return result.value.strip()
